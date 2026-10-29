#!/usr/bin/env python3
"""
gerar_2026.py — painel ao vivo da E2D (substitui a versão antiga com pontos/créditos).

Padrão do teste:
  * margem de 45 min antes de ENTRADA e 25 min antes de SAÍDA/MUDANÇA (tabela de T igual);
  * especiais (JA/SI/CH) fazem exatamente os navios da E12: usa o prático REAL do
    SPPILOTS quando a manobra já tem prático; para as manobras ainda sem prático,
    continua a vez da E12 por simulação;
  * sem pontos nem créditos; restrição Z5 (máx. 1 Z5 por prático regular por dia operacional).

Entradas (mesma pasta deste arquivo):
  inputs/Pauta_de_Manobra_2026.xlsx  — montada pelo ep_atualizar.py (histórico + pauta do SPPILOTS)
  historico_2026.csv                 — histórico do SPPILOTS (com o prático real)
  pauta_rt.csv                       — pauta atual do SPPILOTS (favorito "EP Salvar Pauta")
  inputs/bigramas_2026.xlsx          — escala 2026 (aba RESUMO: dia, D1/D2, P1..P12)
Saída:
  ao_vivo/E2D_2026_v6.xlsx             — lida pelo ep_atualizar.py para montar a página
"""
import sys, os, pickle, subprocess, itertools, time, datetime as dt
from pathlib import Path
from collections import defaultdict
import pandas as pd
import openpyxl, xlsxwriter

try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass
DIR   = Path(__file__).resolve().parent
MOTOR = DIR / 'motor_e2d'
sys.path.insert(0, str(MOTOR))
ESP = {'JA', 'SI', 'CH'}
MARGEM_ENTRADA = int(os.environ.get('MARGEM_ENTRADA', '45'))
MARGEM_OUTRAS  = int(os.environ.get('MARGEM_OUTRAS', '25'))
PAUTA_IN = Path(os.environ.get('PAUTA_IN', DIR / 'inputs' / 'Pauta_de_Manobra_2026.xlsx'))
HIST     = Path(os.environ.get('HIST_IN', DIR / 'historico_2026.csv'))
RT       = Path(os.environ.get('RT_IN', DIR / 'pauta_rt.csv'))
ESCALA   = Path(os.environ.get('ESCALA_IN', DIR / 'inputs' / 'bigramas_2026.xlsx'))
SAIDA    = Path(os.environ.get('SAIDA', DIR / 'ao_vivo' / 'E2D_2026_v6.xlsx'))
t0 = time.time()
J = lambda x: 'JA' if str(x).strip().upper() in ('JÁ', 'JA') else str(x).strip().upper()


def dia_op(t):
    return (t - pd.Timedelta(hours=7, minutes=30)).date()


# ── 1. Pauta: histórico do SPPILOTS (já executado) + pauta atual (daqui para frente) ──
def _ler_csv(path):
    if not path.exists():
        return pd.DataFrame()
    d = pd.read_csv(path, encoding='utf-8-sig')
    d['DT'] = pd.to_datetime(d['datetime'], errors='coerce')
    d = d.dropna(subset=['DT']).copy()
    d['Vessel'] = d['Vessel'].astype(str).str.strip()
    d['Pratico'] = d['Pratico'].fillna('').astype(str) if 'Pratico' in d else ''
    return d
h = _ler_csv(HIST); rt = _ler_csv(RT)


def _completar_dois_praticos(d):
    """Navio de 2 práticos que veio com uma linha só (ainda sem práticos definidos na
    pauta): duplica a linha para que entrem as duas manobras."""
    if d.empty or 'dois_prat' not in d:
        return d
    dois = d['dois_prat'].astype(str).str.strip().str.lower().isin(['true', '1'])
    chave = ['DT', 'Vessel', 'ES'] if 'ES' in d else ['DT', 'Vessel']
    n = d.groupby(chave)['DT'].transform('size')
    falta = d[dois & (n == 1)]
    if len(falta):
        print(f'[pauta] {len(falta)} navio(s) de 2 práticos com uma linha só: duplicados')
        d = pd.concat([d, falta], ignore_index=True).sort_values('DT', kind='stable').reset_index(drop=True)
    return d
h = _completar_dois_praticos(h); rt = _completar_dois_praticos(rt)
if not h.empty:
    # histórico vale até a última manobra raspada; a pauta atual só entra depois disso
    # (se o pauta_rt.csv estiver velho, ele simplesmente não acrescenta nada)
    corte = h.DT.max() + pd.Timedelta(minutes=1)
    if rt.empty or rt.DT.max() < corte:
        print('[AVISO] pauta_rt.csv sem manobras futuras: clique no favorito "EP Salvar Pauta" com o ep_server.py aberto')
    # manobras já passadas que estão na pauta atual mas ainda não entraram no histórico raspado
    # (antes ficavam de fora até a próxima raspagem): entram pela pauta atual
    _orf = rt[(rt.DT < corte) & (rt['Status'].astype(str).str.upper() != 'FUTURO')] if not rt.empty and 'Status' in rt else rt.iloc[0:0]
    if len(_orf):
        _hk = {}
        for _r in h.itertuples():
            _hk.setdefault((_r.Vessel, str(_r.From).strip(), str(_r.To).strip()), []).append(_r.DT)
        _ok = [not any(abs((_r.DT - _t).total_seconds()) <= 6 * 3600
                       for _t in _hk.get((_r.Vessel, str(_r.From).strip(), str(_r.To).strip()), []))
               for _r in _orf.itertuples()]
        _orf = _orf[_ok]
        if len(_orf):
            print(f'[pauta] {len(_orf)} manobra(s) já feitas ainda fora do histórico: mantidas pela pauta atual')
    partes = [h[h.DT < corte]] + ([_orf] if len(_orf) else []) + ([rt[rt.DT >= corte]] if not rt.empty else [])
    p = pd.concat(partes, ignore_index=True)
    p['De'] = p['From']; p['Para'] = p['To']
    p['Termino'] = p['Termino'] if 'Termino' in p else ''
else:   # sem histórico: usa a pauta montada pelo ep_atualizar.py
    p = pd.read_excel(PAUTA_IN).rename(columns={'From': 'De', 'To': 'Para'})
    p['DT'] = pd.to_datetime(p['Data'].astype(str).str.strip() + ' ' + p['Hora'].astype(str).str.strip(),
                             format='%d/%m/%Y %H:%M', errors='coerce')
    p['Vessel'] = p['Vessel'].astype(str).str.strip()
p = p.dropna(subset=['DT']).sort_values('DT', kind='stable').reset_index(drop=True)
p['De'] = p['De'].astype(str).str.strip(); p['Para'] = p['Para'].astype(str).str.strip()
p['Termino'] = p['Termino'].where(p['Termino'].notna(), '').astype(str).replace('nan', '')
print(f'[pauta] {len(p)} manobras, {p.DT.min():%d/%m/%Y} a {p.DT.max():%d/%m/%Y %H:%M}')

# ── 2. Prático real (SPPILOTS): histórico + pauta atual ─────────────────────────
real = {}          # (DT, navio) -> [práticos]
seq_dia = defaultdict(list)   # dia operacional -> ordem real de chamada (E12)
for src in (h, rt):          # pauta atual sobrescreve o histórico
    if src.empty:
        continue
    for k, g in src.groupby(['DT', 'Vessel'], sort=False):
        ps = [J(x) for s in g.Pratico for x in str(s).replace('/', ' ').split() if x and x.upper() not in ('NAN', 'NONE')]
        if ps:
            real[k] = ps
for (t, v), ps in sorted(real.items(), key=lambda kv: kv[0][0]):
    for x in ps:
        if x not in seq_dia[dia_op(t)]:
            seq_dia[dia_op(t)].append(x)
print(f'[sppilots] prático real conhecido em {len(real)} navios')

# ── 3. Escala por dia (planejada + correção pelo que aconteceu de fato) ─────────
R = pd.read_excel(ESCALA, sheet_name='RESUMO')
d0 = dt.date(2026, 1, 1)
PP = ['P%d' % i for i in range(1, 13)]
plano = {}
for i, r in R.iterrows():
    plano[d0 + dt.timedelta(days=i)] = (str(r['TIPO']).strip(), [J(r[c]) if isinstance(r[c], str) else None for c in PP])

# Escala gráfica do SPPilots (lida pelo ep_atualizar.py): quem está realmente de serviço.
# Substitui a escala planejada slot a slot; slot sem leitura mantém o planejado.
_EG = DIR / 'inputs' / 'escala_grafica.json'
if _EG.exists():
    import json as _json
    _n = 0
    for _k, _w in _json.loads(_EG.read_text(encoding='utf-8')).items():
        _d = dt.date.fromisoformat(_k)
        if _d in plano and _w and len(_w) == 12 and sum(1 for x in _w if x) >= 10:
            _tipo, _plan = plano[_d]
            plano[_d] = (_tipo, [J(x) if x else _plan[i] for i, x in enumerate(_w)])
            _n += 1
    print(f'[escala] escala gráfica do SPPilots aplicada em {_n} dias')


def _custo(esc, seq):
    reg = [x for x in esc if x not in ESP and x is not None]
    s = [x for x in seq if x in reg]
    best = None
    for k in range(len(reg)):
        rot = reg[k:] + reg[:k]; ix = {x: i for i, x in enumerate(rot)}
        v = [ix[x] for x in s]
        inv = sum(1 for a in range(len(v)) for b in range(a + 1, len(v)) if v[a] > v[b])
        best = inv if best is None or inv < best else best
    return best or 0


def corrigir(plan, seq):
    """Quem manobrou e não estava na escala entra no lugar de quem não manobrou,
    na posição que deixa a vez real da E12 em ordem (mesma regra validada em 17/10)."""
    W = set(seq)
    outs = [i for i, x in enumerate(plan) if x not in W and x not in ESP]
    ins = [x for x in seq if x not in plan and x not in ESP]
    if not ins or not outs:
        return plan
    ins = ins[:len(outs)]
    melhor = None
    for slots in itertools.permutations(outs, len(ins)):
        e = list(plan)
        for x, s in zip(ins, slots):
            e[s] = x
        c = _custo(e, seq)
        if melhor is None or c < melhor[0]:
            melhor = (c, e)
    return melhor[1]


escala = {}
for d, (tipo, plan) in plano.items():
    escala[d] = (tipo, corrigir(plan, seq_dia[d]) if seq_dia.get(d) else plan)

# turnos D1+D2
dias = sorted(escala)
turnos = []
for i, d in enumerate(dias):
    if escala[d][0] != 'D1':
        continue
    d2 = d + dt.timedelta(days=1)
    nxt = next((x for x in dias[i + 1:] if escala[x][0] == 'D1'), d + dt.timedelta(days=2))
    dia1 = escala[d][1]; dia2 = escala.get(d2, escala[d])[1]
    turnos.append(dict(dt_ini=pd.Timestamp(d) + pd.Timedelta(hours=7.5), dt_d2=pd.Timestamp(d2) + pd.Timedelta(hours=7.5),
                       dt_fim=pd.Timestamp(nxt) + pd.Timedelta(hours=7.5), dia1=dia1, dia2=dia2, Y=list(dia1)))
fim_pauta = p.DT.max()
turnos = [t for t in turnos if t['dt_ini'] <= fim_pauta]
# 1º turno: a fila começa por quem fez a 1ª manobra real do dia
s0 = seq_dia.get(turnos[0]['dt_ini'].date(), [])
if s0 and s0[0] in turnos[0]['dia1']:
    k = turnos[0]['dia1'].index(s0[0]); turnos[0]['Y'] = turnos[0]['dia1'][k:] + turnos[0]['dia1'][:k]
TURNOS = MOTOR / '_turnos_live.pkl'
pickle.dump(turnos, open(TURNOS, 'wb'))
print(f'[escala] {len(turnos)} turnos')

# ── 4. Especiais: E12 real onde há prático; vez da E12 simulada no resto ────────
from ep_algoritmo import gT, BERCO_ZONE
import ep_algoritmo
_Tf = ep_algoritmo._achar_arquivo.__globals__  # só para garantir import
T_df = pd.read_excel(MOTOR / 'inputs' / 'Tempo_T__2025_.xlsx', sheet_name='T Entrada ou saída')
Tmin = {str(r['Berço']).strip(): float(r['T - TEMPO (min)']) for _, r in T_df.iterrows()
        if pd.notna(r['Berço']) and pd.notna(r['T - TEMPO (min)'])}
M_df = pd.read_excel(MOTOR / 'inputs' / 'Tempo_T__2025_.xlsx', sheet_name='Mudanças de entre berços').dropna(subset=['De', 'Para', 'Tm (min)'])
Tmud = {(str(r['De']).strip(), str(r['Para']).strip()): float(r['Tm (min)']) for _, r in M_df.iterrows()}
z5 = lambda a, b: BERCO_ZONE.get(a) == 5 or BERCO_ZONE.get(b) == 5


def fim_de(i):
    t = str(p.Termino.iat[i])
    if ':' in t:
        hh, mm = map(int, t.split(':')[:2]); f = p.DT.iat[i].normalize() + pd.Timedelta(hours=hh, minutes=mm)
        return f if f >= p.DT.iat[i] else f + pd.Timedelta(days=1)
    return p.DT.iat[i] + pd.Timedelta(minutes=gT(p.De.iat[i], p.Para.iat[i], Tmin, Tmud))


FIM = [fim_de(i) for i in range(len(p))]
p['dop'] = [dia_op(t) for t in p.DT]
N = 12
livre = {}; z5c = {}; pred = [None] * len(p); ptr = None; buf = []
for dop, ix in p.groupby('dop', sort=False).indices.items():
    tipo, P = escala.get(dop, ('D2', [None] * 12))
    bold = [x for x in seq_dia.get(dop, []) if x in P][:1] if tipo == 'D1' else []
    if ptr is None or bold:
        ptr = P.index(bold[0]) if bold else 0; buf = []
    ix = list(ix); j = 0
    while j < len(ix):
        n = 2 if j + 1 < len(ix) and p.Vessel.iat[ix[j + 1]] == p.Vessel.iat[ix[j]] and p.DT.iat[ix[j + 1]] == p.DT.iat[ix[j]] else 1
        i0 = ix[j]; t_ = p.DT.iat[i0]; zz = z5(p.De.iat[i0], p.Para.iat[i0])
        conhecido = real.get((t_, p.Vessel.iat[i0]))
        if conhecido:     # E12 real: segue o SPPILOTS e atualiza a vez
            for q, nm in zip(range(n), conhecido + conhecido[:1]):
                pred[ix[j + q]] = nm; livre[nm] = FIM[ix[j + q]]
                if zz: z5c[(dop, nm)] = z5c.get((dop, nm), 0) + 1
                if nm in P:
                    k_ = P.index(nm)
                    if k_ in buf: buf.remove(k_)
                    else:
                        while ptr != k_:
                            buf.append(ptr); ptr = (ptr + 1) % N
                        ptr = (ptr + 1) % N
            j += n; continue
        ok = lambda q: P[q] is not None and livre.get(P[q], pd.Timestamp.min) <= t_ and not (zz and P[q] not in ESP and z5c.get((dop, P[q]), 0) >= 1)
        took = []; k = 0; cls = None; guard = 0
        while len(took) < n and guard < 200:
            guard += 1
            if k >= len(buf):
                if len(buf) + len(took) >= N:
                    cand = [x for x in buf if cls is None or (P[x] in ESP) == cls] or buf
                    livres = [x for x in cand if ok(x)]
                    q = livres[0] if livres else min(cand, key=lambda x: (livre.get(P[x], pd.Timestamp.min), buf.index(x)))
                    took.append(q); buf.remove(q)
                    if cls is None: cls = P[q] in ESP
                    continue
                buf.append(ptr); ptr = (ptr + 1) % N
            q = buf[k]
            if ok(q) and (cls is None or (P[q] in ESP) == cls):
                if cls is None: cls = P[q] in ESP
                took.append(buf.pop(k))
            else:
                k += 1
        for q in took:
            nm = P[q]; pred[ix[j]] = nm; livre[nm] = FIM[ix[j]]
            if zz: z5c[(dop, nm)] = z5c.get((dop, nm), 0) + 1
            j += 1
        if not took:
            j += 1
n_esp = sum(1 for x in pred if x in ESP)
print(f'[especiais] {n_esp} manobras de especiais ({sum(1 for i,x in enumerate(pred) if x in ESP and (p.DT.iat[i],p.Vessel.iat[i]) in real)} pelo SPPILOTS)')

# Giro sem atracar (FB → FB, ex.: "ENTRADA GIRO E SAÍDA DO TECON"): T pela linha "GIRO <área>" da tabela
# de T, conforme a área citada na pauta; sem área identificada, usa GIRO TECON. (Regra de 27/10/2026.)
_GIRO_T = {}
try:
    _tt = pd.read_excel(DIR / 'inputs' / 'Tempo_T__2025_.xlsx', sheet_name='T Entrada ou saída')
    _gt = {str(b).strip().upper(): float(t) for b, t in zip(_tt['Berço'], _tt['T - TEMPO (min)'])
           if str(b).strip().upper().startswith('GIRO') and pd.notna(t)}
    _areas = [('GIRO EMBRAPORT', ('EMBRAPORT', 'DP WORLD', 'DPW')), ('GIRO BTP', ('BTP',)),
              ('GIRO ALEMOA', ('ALEMOA', 'LEMOA')), ('GIRO TECON', ('TECON', 'TCOM', 'SANTOS BRASIL'))]
    _cols = [c for c in ('OBS', 'Obs', 'Agente', 'Celulas') if c in p.columns]
    for _i in range(len(p)):
        if p.De.iat[_i] != 'FB' or p.Para.iat[_i] != 'FB':
            continue
        _txt = ' '.join(str(p[c].iat[_i]) for c in _cols).upper()
        _area = next((a for a, kws in _areas if any(k in _txt for k in kws)), 'GIRO TECON')
        if _area in _gt:
            _GIRO_T[(p.DT.iat[_i].strftime('%Y-%m-%d %H:%M'), 'FB', 'FB')] = _gt[_area]
    if _GIRO_T:
        print(f'[giro] {len(_GIRO_T)} manobra(s) de giro com T da tabela')
except Exception as _e:
    print(f'[giro] não aplicado: {_e}')
pickle.dump(_GIRO_T, open(MOTOR / '_giro_T.pkl', 'wb'))

# arquivos de entrada do motor
PAUTA_M = MOTOR / '_pauta_live.xlsx'
pd.DataFrame({'Data': p.DT.dt.strftime('%d/%m/%Y'), 'Hora': p.DT.dt.strftime('%H:%M'), 'Termino': p.Termino.astype(str),
              'Vessel': p.Vessel, 'From': p.De, 'To': p.Para}).to_excel(PAUTA_M, index=False)
ESP_F = MOTOR / '_especiais_live.xlsx'
wb = xlsxwriter.Workbook(str(ESP_F)); ws = wb.add_worksheet()
ws.write_row(0, 0, ['Seq', 'Navio', 'Pratico'])
for i in range(len(p)):
    ws.write_row(i + 1, 0, [i + 1, p.Vessel.iat[i], pred[i] or ''])
wb.close()

# ── 5. Motor E2D ────────────────────────────────────────────────────────────────
# Escalação congelada (regra de 29/10/2026): manobra que já COMEÇOU (linha azul na pauta do SPPilots:
# PROXIMO/CONFIRMADO) ou TERMINOU (linha vermelha: AGENDADO/EXECUTADO, ou já no histórico) não muda mais.
# O prático que a E2D tinha nessa hora fica fixo. Linhas brancas e lilases ainda podem mudar (inclusive com atraso).
# Só regulares; manobra de especial nunca é congelada (o motor ignora, ep_sim_v3_corrigido).
STATUS_FIXO = {'PROXIMO', 'CONFIRMADO', 'AGENDADO', 'EXECUTADO'}
_mn = lambda t: pd.Timestamp(t).floor('min')
_fixas = set()
if not h.empty:
    _fixas |= {(_mn(t), str(v).strip()) for t, v in zip(h.DT, h.Vessel)}
if not rt.empty and 'Status' in rt:
    _fixas |= {(_mn(t), str(v).strip()) for t, v, s_ in zip(rt.DT, rt.Vessel, rt.Status)
               if str(s_).strip().upper() in STATUS_FIXO}
CONG = DIR / 'dados' / 'escalacao_finalizada.csv'
# (29/10/2026) O congelamento anterior (1 h antes, gerado com a troca direta) foi DESCARTADO a pedido da Praticagem:
# não se congela o que foi gerado errado. O painel foi refeito do zero com as regras corretas e o congelamento
# vale daí em diante. dados/escalacao_congelada.csv fica guardado, sem uso.
_cg = pd.read_csv(CONG, dtype=str) if CONG.exists() else pd.DataFrame(columns=['DT', 'Vessel', 'Pratico', 'registrado_em'])
_fz = defaultdict(list)
for _r in _cg.itertuples():
    _d = pd.Timestamp(_r.DT)
    _fz[(_d.strftime('%d/%m/%Y'), _d.strftime('%H:%M'), str(_r.Vessel).strip())].append(str(_r.Pratico).strip())
pickle.dump(dict(_fz), open(MOTOR / '_congeladas.pkl', 'wb'))
OUT_M = '_E2D_live.xlsx'
env = dict(os.environ, PAUTA_FILE=str(PAUTA_M), ESP_FONTE=str(ESP_F), TURNOS=str(TURNOS),
           MARGEM_ENTRADA=str(MARGEM_ENTRADA), MARGEM_OUTRAS=str(MARGEM_OUTRAS), E2D_OUT=OUT_M,
           FB_OUT=str(MOTOR / '_fallback.pkl'), PYTHONIOENCODING='utf-8', BTP_OV=str(MOTOR / '_giro_T.pkl'), FROZEN=str(MOTOR / '_congeladas.pkl'))
r = subprocess.run([sys.executable, 'motor_e2d.py'], cwd=str(MOTOR), env=env, capture_output=True,
                   text=True, encoding='utf-8', errors='replace')
if r.returncode != 0:
    print(r.stdout[-1500:]); print(r.stderr[-3000:]); sys.exit(1)
print('[motor]', ' | '.join(l for l in r.stdout.splitlines() if 'Etapa' in l or 'Especiais' in l or 'Fallback' in l or 'congeladas' in l))

# grava as manobras que já começaram ou terminaram (azul/vermelha) e ainda não estavam congeladas
try:
    _agora = pd.Timestamp(dt.datetime.now())
    _ja = {(_mn(a), str(b).strip()) for a, b in zip(_cg.DT, _cg.Vessel)}
    _novos = defaultdict(list)
    for _row in openpyxl.load_workbook(MOTOR / OUT_M, read_only=True).worksheets[0].iter_rows(min_row=2, values_only=True):
        if not isinstance(_row[0], int):
            continue
        _d = pd.to_datetime(f'{_row[1]} {_row[2]}', dayfirst=True)
        _pr, _nv = str(_row[4] or '').strip(), str(_row[5] or '').strip()
        if (_mn(_d), _nv) in _fixas and (_mn(_d), _nv) not in _ja and _pr and _pr not in ESP:
            _novos[(_d, _nv)].append(_pr)
    if _novos:
        _add = pd.DataFrame([{'DT': k[0].strftime('%Y-%m-%d %H:%M'), 'Vessel': k[1], 'Pratico': p_,
                              'registrado_em': _agora.strftime('%Y-%m-%d %H:%M')}
                             for k, v in _novos.items() for p_ in v])
        CONG.parent.mkdir(parents=True, exist_ok=True)
        pd.concat([_cg, _add], ignore_index=True).to_csv(CONG, index=False, encoding='utf-8')
        print(f'[congela] {len(_novos)} manobra(s) começada(s)/terminada(s) agora: prático fixo')
except Exception as _e:
    print(f'[congela] não gravado: {_e}')

# ── 6. Planilha no formato que o ep_atualizar.py lê ─────────────────────────────
src = openpyxl.load_workbook(MOTOR / OUT_M, read_only=True).worksheets[0]
SAIDA.parent.mkdir(parents=True, exist_ok=True)
wb = xlsxwriter.Workbook(str(SAIDA)); ws = wb.add_worksheet('E2D 2026')
ws.write(0, 0, f'E2D 2026 — margem {MARGEM_ENTRADA} min antes de entrada / {MARGEM_OUTRAS} min antes de saída e mudança — gerado {dt.datetime.now():%d/%m/%Y %H:%M}')
ws.write_row(1, 0, ['Seq', 'Data', 'Hora', 'Termino', 'Pratico', 'Navio', 'De', 'Para', 'Pts(cont)', 'Pts(real)', 'Total', 'Status', '#Pos', 'Escalacao'])
ri = 2; n = 0
for row in src.iter_rows(min_row=2, values_only=True):
    if not isinstance(row[0], int):
        continue
    seq, data, hora, term, prat, nav, de, para, st, esc = row[:10]
    ws.write_row(ri, 0, [seq, data, hora, term or '', prat, nav, de, para, '', '', '', st or '', '', esc or ''])
    ri += 1; n += 1
wb.close()
print(f'[ok] {SAIDA} — {n} manobras — {time.time()-t0:.0f}s')
